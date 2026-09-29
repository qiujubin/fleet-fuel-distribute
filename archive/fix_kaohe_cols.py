# -*- coding: utf-8 -*-
"""补齐 09-19 写入考核表的 11 行的 21~24 列 (U/V/W/X)"""
import os
import openpyxl

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'

# sid -> (sheet, 模板文件名)
targets = {
    '260918000408': ('每日油耗数据', 'D14-4.2米-粤BT802F.xlsx'),
    '260918000409': ('每日油耗数据', 'B13-9.6米-粤BNT993.xlsx'),
    '260918000413': ('每日油耗数据', 'D17-4.2米-京LPW138.xlsx'),
    '260918000415': ('每日油耗数据', 'D14-4.2米-粤BT802F.xlsx'),
    '260919000416': ('每日油耗数据', 'C2-7.6米-粤BLY079.xlsx'),
    '260919000417': ('每日油耗数据', 'B15-9.6米-粤BNJ052.xlsx'),
    '260919000419': ('每日油耗数据', 'D12-4.2米-粤BN2S63.xlsx'),
    '260919000420': ('每日油耗数据', 'D15-4.2米-粤BY2J27.xlsx'),
    '260919000423': ('每日油耗数据', 'D21-4.2米-粤B17K86.xlsx'),
    '260919000424': ('每日油耗数据', 'B13-9.6米-粤BNT993.xlsx'),
    '260918000410': ('每日尿素数据', '尿素B16-9.6米-粤BNV030.xlsx'),
}

# 读备注
notes = {}
for sid, (_, fn) in targets.items():
    wb = openpyxl.load_workbook(os.path.join(TPL, fn), read_only=True, data_only=True)
    for row in wb['加油数据'].iter_rows(min_row=2, values_only=True):
        if row and str(row[0] or '').strip() == sid:
            notes[sid] = str(row[13] or '').strip()
            break
    wb.close()

wbk = openpyxl.load_workbook(KAOHE, keep_vba=True)
for sid, (sheet, _) in targets.items():
    ws = wbk[sheet]
    row = None
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or '').strip() == sid:
            row = r; break
    assert row, sid
    I = float(ws.cell(row, 9).value or 0)
    Q = float(ws.cell(row, 17).value or 0)
    R = ws.cell(row, 18).value
    T = ws.cell(row, 20).value
    G = ws.cell(row, 7).value
    U = round(I / Q, 10) if Q else None
    V = round((float(R) + float(T)) / 2, 10) if (R is not None and T is not None) else None
    W = G.date() if hasattr(G, 'date') else G
    X = notes.get(sid) or 0
    ws.cell(row, 21, U)
    ws.cell(row, 22, V)
    c23 = ws.cell(row, 23, W); c23.number_format = 'yyyy-mm-dd'
    ws.cell(row, 24, X)
    print(f'{sid} 行{row}: U={U} V={V} W={W} X={X!r}')
wbk.save(KAOHE)
print('考核表 21~24 列已补齐保存')
