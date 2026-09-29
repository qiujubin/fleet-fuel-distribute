# -*- coding: utf-8 -*-
"""对账 v2: 按 文件的油/尿素属性 分流, 只统计考核文件起点(2025-10-05)之后的"是"行"""
import os, glob, datetime
import openpyxl
from collections import defaultdict

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
CUTOFF = '2025-10-05'

def gs(wb):
    for n in ('计算模板', '油耗计算'):
        if n in wb.sheetnames: return wb[n]

shi = {'油': defaultdict(set), '尿素': defaultdict(set)}   # tag -> plate -> {(g19,sid)}
pre = {'油': 0, '尿素': 0}
for fp in sorted(glob.glob(os.path.join(TPL, '*.xlsx'))):
    fn = os.path.basename(fp)
    if fn.startswith('~$') or fn.startswith('模板'): continue
    tag = '尿素' if fn.startswith('尿素') else '油'
    plate = os.path.splitext(fn)[0].split('-')[-1].strip().upper()
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    wst = gs(wb); wsd = wb['加油数据']
    if wst is None: wb.close(); continue
    rec_full = {}
    for row in wsd.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None:
            rec_full[str(row[12])[:19]] = str(row[8] or '').strip()
    for row in wst.iter_rows(min_row=2, values_only=True):
        if row and row[6] is not None:
            g19 = str(row[6])[:19]
            if rec_full.get(g19) == '是':
                if g19[:10] < CUTOFF:
                    pre[tag] += 1
                else:
                    shi[tag][plate].add((g19, str(row[0] or '').strip()))
    wb.close()

wbk = openpyxl.load_workbook(KAOHE, data_only=True)
kao = {'油': defaultdict(set), '尿素': defaultdict(set)}
for sn, tag in (('每日油耗数据', '油'), ('每日尿素数据', '尿素')):
    for row in wbk[sn].iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, '', '0'): continue
        plate = str(row[3] or '').strip().upper()
        kao[tag][plate].add((str(row[6])[:19] if row[6] else '', str(row[0]).strip()))

print(f'考核起点({CUTOFF})之后的"是"行: 油{sum(len(v) for v in shi["油"].values())} 尿素{sum(len(v) for v in shi["尿素"].values())}')
print(f'起点之前的"是"行(考核文件尚不存在, 不计): 油{pre["油"]} 尿素{pre["尿素"]}')
print()
tot_m = tot_e = 0
for tag in ('油', '尿素'):
    plates = sorted(set(list(shi[tag].keys()) + list(kao[tag].keys())))
    for plate in plates:
        v = shi[tag].get(plate, set())
        k = kao[tag].get(plate, set())
        miss = v - k; extra = k - v
        if miss or extra:
            tot_m += len(miss); tot_e += len(extra)
            print(f'[{tag}] {plate}: 模板是行{len(v)} 考核{len(k)} | 考核缺{len(miss)} 考核多{len(extra)}')
            for g, s in sorted(miss)[:4]: print(f'      缺 {g} {s}')
            for g, s in sorted(extra)[:4]: print(f'      多 {g} {s}')
print(f'\n真实差异: 考核缺 {tot_m} 条, 考核孤儿 {tot_e} 条')
