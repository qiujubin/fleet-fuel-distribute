# -*- coding: utf-8 -*-
"""按有效"是"规则重新写入 09-19 的考核行 (10条: 9油+1尿素, D14只写415累计行)"""
import os, datetime
import openpyxl

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
DT_FMT = 'yyyy-mm-dd hh:mm:ss'

# (sid, 模板文件, 模板行r, F时间或None=首周期, cycle行列表)
items = [
    ('260918000409', 'B13-9.6米-粤BNT993.xlsx', 148, '2026-09-17 17:27:56', []),
    ('260919000424', 'B13-9.6米-粤BNT993.xlsx', 149, '2026-09-18 15:18:20', []),
    ('260918000413', 'D17-4.2米-京LPW138.xlsx', 200, '2026-09-17 09:27:15', []),
    ('260918000415', 'D14-4.2米-粤BT802F.xlsx', 198, '2026-09-16 14:11:57', [197]),
    ('260919000416', 'C2-7.6米-粤BLY079.xlsx', 152, '2026-09-17 08:30:08', []),
    ('260919000417', 'B15-9.6米-粤BNJ052.xlsx', 191, '2026-09-16 23:35:53', [190]),
    ('260919000419', 'D12-4.2米-粤BN2S63.xlsx', 157, '2026-09-16 13:11:14', []),
    ('260919000420', 'D15-4.2米-粤BY2J27.xlsx', 144, '2026-09-17 14:25:51', []),
    ('260919000423', 'D21-4.2米-粤B17K86.xlsx', 225, '2026-09-17 10:34:12', []),
    ('260918000410', '尿素B16-9.6米-粤BNV030.xlsx', 85, '2026-09-17 01:42:32', []),
]

def num(v):
    try: return float(v)
    except: return None

def rec_by_time(wsd, plate, t):
    for r in range(2, wsd.max_row + 1):
        mt = wsd.cell(r, 13).value
        if mt and str(mt)[:19] == str(t)[:19] and \
           str(wsd.cell(r, 2).value or '').strip().upper() == plate:
            return {'mile': num(wsd.cell(r, 4).value), 'gps': num(wsd.cell(r, 5).value),
                    'amt': num(wsd.cell(r, 6).value), 'vol': num(wsd.cell(r, 7).value),
                    'note': str(wsd.cell(r, 14) or '').strip()}
    return None

wbk = openpyxl.load_workbook(KAOHE, keep_vba=True)
try:
    for sid, fn, r, F_time, cycle_rows in items:
        is_urea = fn.startswith('尿素')
        wbn = openpyxl.load_workbook(os.path.join(TPL, fn), data_only=True)
        wstv = None
        for n in ('计算模板', '油耗计算'):
            if n in wbn.sheetnames: wstv = wbn[n]
        wsd = wbn['加油数据']
        plate = str(wstv.cell(2, 4).value or '').strip()
        type_v = str(wstv.cell(2, 5).value or '')
        recG = rec_by_time(wsd, plate, wstv.cell(r, 7).value)
        recF = rec_by_time(wsd, plate, F_time)
        amt = recG['amt']; vol = recG['vol']
        for x in cycle_rows:
            gt = wstv.cell(x, 7).value
            ri = rec_by_time(wsd, plate, gt) if gt else None
            if ri:
                amt = (amt or 0) + (ri['amt'] or 0)
                vol = (vol or 0) + (ri['vol'] or 0)
        L, M = recF['gps'], recF['mile']
        N, O = recG['gps'], recG['mile']
        P = round(N - L, 4) if (L is not None and N is not None) else None
        Q = round(O - M, 4) if (M is not None and O is not None) else None
        Rv = round(vol / P * 100, 6) if P else None
        Sv = round(amt / P, 6) if P else None
        Tv = round(vol / Q * 100, 6) if Q else None
        Uv = round(amt / Q, 10) if Q else None
        Vv = round((Rv + Tv) / 2, 10) if (Rv is not None and Tv is not None) else None
        note = recG['note']
        sheet = '每日尿素数据' if is_urea else '每日油耗数据'
        ws = wbk[sheet]
        # 去重
        exists = any(str(ws.cell(x, 1).value or '').strip() == sid for x in range(2, ws.max_row + 1))
        if exists:
            print(f'跳过(已存在): {sid}'); continue
        nr = 1
        for x in range(1, ws.max_row + 1):
            if any(ws.cell(x, c).value not in (None, '') for c in range(1, 25)): nr = x
        nr += 1
        rowvals = [sid, r - 1, recG and None, plate, type_v, None, None, '是',
                   amt, vol, recG and None, L, M, N, O, P, Q, Rv, Sv, Tv, Uv, Vv, None, None]
        # 登记人/价格/时间从模板取
        person = None; price = None
        for rr in range(2, wsd.max_row + 1):
            if str(wsd.cell(rr, 13).value or '')[:19] == str(wstv.cell(r, 7).value)[:19]:
                person = wsd.cell(rr, 12).value; price = num(wsd.cell(rr, 8).value)
                break
        rowvals[2] = person; rowvals[10] = price
        rowvals[5] = wstv.cell(r, 6).value; rowvals[6] = wstv.cell(r, 7).value
        rowvals[22] = (recG and None) or recG  # placeholder
        rowvals[22] = wstv.cell(r, 7).value.date() if hasattr(wstv.cell(r, 7).value, 'date') else None
        rowvals[23] = note or 0
        for c, v in enumerate(rowvals, 1):
            cell = ws.cell(nr, c, v)
            if c in (6, 7): cell.number_format = DT_FMT
            if c == 23: cell.number_format = 'yyyy-mm-dd'
        print(f'{sid} {plate} -> {sheet} 第{nr}行 (I={amt}, J={vol})')
        wbn.close()
    wbk.save(KAOHE)
    print('考核表重新写入完成并保存')
except PermissionError as e:
    print('文件被占用! 请先关闭 Excel 中打开的考核表, 再重试.', e)
