# -*- coding: utf-8 -*-
import openpyxl, zipfile, os
from collections import Counter
TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
CUR = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
SUM = r'C:\Users\Jubin\Desktop\源数据\油耗分析汇总.xlsm'

for fn, sids in [('B11-9.6米-粤BFE485.xlsx', ['260919000429', '260920000440', '260920000442']),
                 ('D1-4.2米-粤A2K1E1.xlsx', ['260920000441'])]:
    wb = openpyxl.load_workbook(os.path.join(TPL, fn), read_only=True, data_only=True)
    wsd = wb['加油数据']
    sids_in = [str(r[0]) for r in wsd.iter_rows(min_row=2, values_only=True) if r and r[0] is not None]
    dup = [k for k, c in Counter(sids_in).items() if c > 1]
    ok = all(s in sids_in for s in sids) and not dup
    print(fn[:3], '记录', len(sids_in), '| 重复', dup or '无', '| 新增齐全' if ok else '| 异常!')
    wb.close()

wb = openpyxl.load_workbook(CUR, read_only=True, data_only=True)
ws = wb['每日油耗数据']
kao = [str(r[0]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0] not in (None, '', '0')]
dup = [k for k, c in Counter(kao).items() if c > 1]
targets = ['260919000427', '260919000429', '260920000431', '260920000432', '260920000434',
           '260920000435', '260920000436', '260920000437', '260920000438', '260920000439',
           '260920000440', '260920000441']
miss = [t for t in targets if t not in kao]
print('考核表', len(kao), '行 | 重复', dup or '无', '| 今日目标缺失', miss or '无')
ws2 = wb['每日加油记录']
n = sum(1 for r in ws2.iter_rows(min_row=2, values_only=True) if r and r[0])
print('每日加油记录', n, '行')
wb.close()

wb = openpyxl.load_workbook(SUM, read_only=True, data_only=True)
ws = wb['Sheet1']
n = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r and any(c is not None for c in r))
print('汇总', n, '行')
wb.close()

z = zipfile.ZipFile(CUR)
rx = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
print('宏工程+引用:', '完整' if ('xl/JDEData.bin' in z.namelist() and 'jdeExtension' in rx) else '异常!')
z.close()
