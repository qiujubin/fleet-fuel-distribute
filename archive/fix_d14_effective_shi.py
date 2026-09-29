# -*- coding: utf-8 -*-
"""按"同一天只算最后一个有效是"规则修正 D14-粤BT802F 的模板与考核"""
import os, datetime
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
DT_FMT = 'yyyy-mm-dd hh:mm:ss'

p = os.path.join(TPL, 'D14-4.2米-粤BT802F.xlsx')
wb = openpyxl.load_workbook(p, data_only=False)
wbv = openpyxl.load_workbook(p, data_only=True)
wst, wstv = wb['计算模板'], wbv['计算模板']
wsd = wb['加油数据']

# 上一有效"是" = 行196 的 G (2026-09-16 14:11:57)
prevG = wstv.cell(196, 7).value
assert str(prevG)[:19] == '2026-09-16 14:11:57', prevG
# 该时间记录的 GPS/仪表 (考核 L/M 用)
recF = None
for r in range(2, wsd.max_row + 1):
    if str(wsd.cell(r, 13).value or '')[:19] == '2026-09-16 14:11:57':
        recF = (wsd.cell(r, 4).value, wsd.cell(r, 5).value)  # (当前里程, GPS里程)
        break
print('上一有效是的记录: 里程=', recF[0], 'GPS=', recF[1])

# --- 模板修正 ---
# 行197 (408, 当天第一个是): F 清空, I/J 保持 base
wst.cell(197, 6, None)
# 行198 (415, 当天最后是): F = 行196 G, I/J = base + I197 (+J197)
wst.cell(198, 6, prevG).number_format = DT_FMT
base_i = f'=XLOOKUP(D198&G198,加油数据!B:B&加油数据!M:M,加油数据!F:F,0)+I197'
base_j = f'=XLOOKUP(D198&G198,加油数据!B:B&加油数据!M:M,加油数据!G:G,0)+J197'
wst.cell(198, 9, ArrayFormula('I198', base_i))
wst.cell(198, 10, ArrayFormula('J198', base_j))
wb.save(p); wb.close()
print('模板: 行197 F清空; 行198 F=09-16 14:11:57, I/J=base+I197')

# --- 考核修正 ---
wbk = openpyxl.load_workbook(KAOHE, keep_vba=True)
w1 = wbk['每日油耗数据']
# 删除 408 行
r408 = None
for r in range(2, w1.max_row + 1):
    if str(w1.cell(r, 1).value or '').strip() == '260918000408':
        r408 = r; break
assert r408, '408 not found'
w1.delete_rows(r408, 1)
print(f'考核: 已删除 408 行(原第{r408}行)')
# 更新 415 行 (删除后行号-1 -> 3262)
r415 = None
for r in range(2, w1.max_row + 1):
    if str(w1.cell(r, 1).value or '').strip() == '260918000415':
        r415 = r; break
assert r415, '415 not found'
# 数据: 415 自身 + 408 累计
I = 586.55 + 608.0        # 1194.55
J = 73.78 + 76.68         # 150.46
L = float(recF[1]); M = float(recF[0])   # GPS开始/仪表开始
N = 506540.5; O = 344384.0               # 415 的 GPS/仪表 (与之前一致)
P = round(N - L, 4); Q = round(O - M, 4)
Rv = round(J / P * 100, 6); Sv = round(I / P, 6); Tv = round(J / Q * 100, 6)
Uv = round(I / Q, 10); Vv = round((Rv + Tv) / 2, 10)
wbv2 = openpyxl.load_workbook(p, data_only=True)
wsd2 = wbv2['加油数据']
for r in range(2, wsd2.max_row + 1):
    if str(wsd2.cell(r, 13).value or '')[:19] == '2026-09-18 23:02:36':
        N = float(wsd2.cell(r, 5).value); O = float(wsd2.cell(r, 4).value)
        P = round(N - L, 4); Q = round(O - M, 4)
        Rv = round(J / P * 100, 6); Sv = round(I / P, 6); Tv = round(J / Q * 100, 6)
        Uv = round(I / Q, 10); Vv = round((Rv + Tv) / 2, 10)
        break
wbv2.close()
vals = {6: prevG, 9: I, 10: J, 12: L, 13: M, 14: N, 15: O, 16: P, 17: Q,
        18: Rv, 19: Sv, 20: Tv, 21: Uv, 22: Vv}
for c, v in vals.items():
    cell = w1.cell(r415, c, v)
    if c == 6: cell.number_format = DT_FMT
print(f'考核: 415 行{r415} 已改为累计值: F={prevG} I={I} J={J} P={P} Q={Q} R={Rv} T={Tv}')
wbk.save(KAOHE)
print('考核表已保存')
