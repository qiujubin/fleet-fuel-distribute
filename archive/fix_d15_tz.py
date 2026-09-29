# -*- coding: utf-8 -*-
"""修 D15 粤BY2J27: 行151 的 F 锚点被时区 bug 写偏(-8h), 应指向行149 的 G(09-24 23:04)"""
import os, shutil, datetime, traceback
import openpyxl
import pythoncom, win32com.client as wc
from openpyxl.utils.datetime import from_excel
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rebuild_formulas import rebuild

D15 = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板\D15-4.2米-粤BY2J27.xlsx'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
DT_FMT = 'yyyy-mm-dd hh:mm:ss'
BAK = r'E:\WorkSpace\workbuddy\2026-09-18-16-38-15\backup_fix448'
os.makedirs(BAK, exist_ok=True)

# ---- 0) 备份 ----
shutil.copy2(D15, os.path.join(BAK, os.path.basename(D15)))
shutil.copy2(KAOHE, os.path.join(BAK, os.path.basename(KAOHE)))
print('备份完成')

# ---- 1) 修 F: 行151 = 行149 的 G ----
wb = openpyxl.load_workbook(D15, data_only=False)
wst = wb['计算模板']
g149 = wst.cell(149, 7).value
bad_f = wst.cell(151, 6).value
print(f'行149 G = {g149}  ({from_excel(g149) if isinstance(g149,(int,float)) else ""})')
print(f'行151 F(旧, 错误) = {bad_f}  ({from_excel(bad_f) if isinstance(bad_f,(int,float)) else ""})')
assert g149 is not None
wst.cell(151, 6).value = g149
wst.cell(151, 6).number_format = DT_FMT
wb.save(D15)
wb.close()
print('行151 F 已改为行149 的 G')

# ---- 2) 重建公式 ----
n, filled = rebuild(D15)
print(f'rebuild: {n} 公式 / {filled} 行G')

# ---- 3) COM 重算 + 更新考核行 ----
pythoncom.CoInitialize()
app = None
try:
    app = wc.DispatchEx('Ket.Application')
    app.Visible = False
    app.DisplayAlerts = False
    wbv = app.Workbooks.Open(D15, 0, False)
    wsv = wbv.Worksheets('计算模板')
    app.Calculate()
    rng = wsv.Range(wsv.Cells(151, 6), wsv.Cells(151, 22)).Value2
    if not isinstance(rng, tuple):
        rng = (rng,)
    if rng and isinstance(rng[0], tuple):
        rng = rng[0]
    print('行151 F..V:', [round(x, 4) if isinstance(x, float) else x for x in rng])
    wbv.Save(); wbv.Close(False)

    def n_(x):
        try: return float(x)
        except (TypeError, ValueError): return None

    F_s, G_s = n_(rng[0]), n_(rng[1])
    F_t = from_excel(F_s) if F_s else None
    G_t = from_excel(G_s) if G_s else None
    print(f'F={F_t}  G={G_t}')

    wbk = app.Workbooks.Open(KAOHE, 0, False)
    wsd = wbk.Worksheets('每日油耗数据')
    # 定位: 车牌=粤BY2J27 且 时间B(第7列) 的日 = 09-25
    target = None
    for i in range(2, 4000):
        p = str(wsd.Cells(i, 4).Value2 or '').strip()
        if p != '粤BY2J27':
            continue
        tb = n_(wsd.Cells(i, 7).Value2)
        if tb:
            d = from_excel(tb)
            if d.year == 2026 and d.month == 9 and d.day == 25:
                target = i
    if not target:
        print('!!! 未找到考核行')
    else:
        rowvals = [None, None, None, None, None, F_t, G_t, '是',
                   n_(rng[3]), n_(rng[4]), n_(rng[5]), n_(rng[6]), n_(rng[7]), n_(rng[8]), n_(rng[9]),
                   n_(rng[10]), n_(rng[11]), n_(rng[12]), n_(rng[13]), n_(rng[14]),
                   n_(rng[15]), n_(rng[16]), G_t, 0]
        for c in (1, 2, 3, 5):
            rowvals[c - 1] = wsd.Cells(target, c).Value2
        wsd.Cells(target, 1).NumberFormat = '@'
        wsd.Range(wsd.Cells(target, 1), wsd.Cells(target, 24)).Value = tuple([tuple(rowvals)])
        for c in (6, 7):
            wsd.Cells(target, c).NumberFormat = DT_FMT
        wsd.Cells(target, 23).NumberFormat = 'yyyy-mm-dd'
        print(f'考核行{target} 已更新: A={F_t} B={G_t} 金额={rowvals[8]} 升={rowvals[9]} 里程P={rowvals[15]} Q={rowvals[16]}')
        wbk.Save()
    wbk.Close(False)
finally:
    if app is not None:
        try: app.Quit()
        except Exception: pass
    pythoncom.CoUninitialize()
print('完成')
