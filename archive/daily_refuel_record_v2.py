# -*- coding: utf-8 -*-
"""刷新「每日加油记录」sheet
源: 数据文件夹里最新的 车辆加油_管理 导出文件
取: 登记日期 = n-1 (n=运行当天) 的记录
规则: 仅加油(排除尿素/错标尿素的加油记录: 单价<5);
      排除词(发动机/发电机/小油箱/新车)记录 "是" 强制改 "否";
      同车牌同日多条 -> 有"是"留最后一个"是", 全"否"留最后一个。整表替换。
"""
import os, sys, glob, re, datetime, subprocess
import openpyxl
from python_calamine import CalamineWorkbook

DATA_DIR = r'C:\Users\Jubin\Desktop\数据'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
RESTORE = r'E:\WorkSpace\workbuddy\2026-09-18-16-38-15\restore_wps_parts.py'
EXCLUDE_KW = ['发动机', '发电机', '小油箱', '新车']
DT_FMT, D_FMT = 'yyyy-mm-dd hh:mm:ss', 'yyyy-mm-dd'

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return v

def parse_dt(v):
    if isinstance(v, datetime.datetime): return v
    s = str(v).strip()
    for f in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
        try: return datetime.datetime.strptime(s, f)
        except ValueError: pass
    return None

# 1) 找最新导出文件(按文件名日期倒序, 取第一个表头完整的)
files = []
for fp in glob.glob(os.path.join(DATA_DIR, '车辆加油_管理*.xlsx')):
    m = re.search(r'\((\d{4}-\d{2}-\d{2})\)', os.path.basename(fp))
    if m: files.append((m.group(1), fp))
files.sort(reverse=True)
n = datetime.date.today()
n1 = n - datetime.timedelta(days=1)
print(f'今天 n={n}, 取 n-1={n1}')

recs = []
used_file = None
for d, fp in files:
    wb = CalamineWorkbook.from_path(fp)
    rows = wb.get_sheet_by_name(wb.sheet_names[0]).to_python()
    if not rows: continue
    hdr = [str(x).strip() for x in rows[0]]
    if '系统编号' not in hdr: continue
    used_file = os.path.basename(fp)
    for row in rows[1:]:
        r = dict(zip(hdr, row))
        sid = str(r.get('系统编号') or '').strip()
        if not sid: continue
        t = parse_dt(r.get('登记时间'))
        dd = parse_dt(r.get('登记日期'))
        if dd is None or dd.date() != n1: continue
        recs.append({
            'sid': sid, 'plate': str(r.get('车牌号码') or '').strip(),
            'date': dd, 'mile': num(r.get('当前里程')), 'gps': num(r.get('GPS里程')),
            'amt': num(r.get('加注金额(元)')), 'vol': num(r.get('加注油量（升）')),
            'price': num(r.get('加注单价（元）')), 'full': str(r.get('是否加满油') or '').strip(),
            'type': str(r.get('加油类型') or '').strip(), 'loc': str(r.get('当前位置') or '').strip(),
            'person': str(r.get('登记人') or '').strip(), 'time': t,
            'note': str(r.get('备注信息') or '').strip()})
    break
print('源文件:', used_file, '| n-1 记录(原始):', len(recs))

# 2) 过滤 + 规则
pool = []
for r in recs:
    if '尿素' in r['type']: continue
    p = r['price']
    if isinstance(p, float) and 0 < p < 5: continue    # 错标成加油的尿素
    full = r['full']
    if any(k in r['note'] for k in EXCLUDE_KW): full = '否'
    r['full'] = full
    pool.append(r)
by_plate = {}
for r in pool:
    by_plate.setdefault(r['plate'].upper(), []).append(r)
keep = []
for plate, items in by_plate.items():
    items.sort(key=lambda x: str(x['time'])[:19])
    shi = [x for x in items if x['full'] == '是']
    keep.append(shi[-1] if shi else items[-1])
keep.sort(key=lambda x: str(x['time'])[:19])
print(f'过滤后候选 {len(pool)} 条, 逐车保留后 {len(keep)} 条')

# 3) 重写 sheet
wbk = openpyxl.load_workbook(KAOHE, keep_vba=True)
ws = wbk['每日加油记录']
last = 1
for r in range(1, ws.max_row + 1):
    if any(ws.cell(r, c).value not in (None, '') for c in range(1, 15)): last = r
for r in range(2, last + 1):
    for c in range(1, 15): ws.cell(r, c, None)
for i, x in enumerate(keep, start=2):
    vals = [x['sid'], x['plate'], x['date'], x['mile'], x['gps'], x['amt'], x['vol'], x['price'],
            x['full'], x['type'], x['loc'], x['person'], x['time'], x['note']]
    for c, v in enumerate(vals, start=1):
        ws.cell(i, c, v)
    ws.cell(i, 3).number_format = D_FMT
    ws.cell(i, 13).number_format = DT_FMT
    print(f"  行{i}: {x['sid']} {x['plate']} {x['time']} 加满[{x['full']}] ¥{x['amt']} {x['vol']}L 备注[{x['note'] or '-'}]")
wbk.save(KAOHE)
print('每日加油记录已刷新')
subprocess.run([sys.executable, RESTORE], check=False)
print('WPS部件已恢复')
