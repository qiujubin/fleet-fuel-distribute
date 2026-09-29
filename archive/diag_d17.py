# -*- coding: utf-8 -*-
import openpyxl, os
from openpyxl.worksheet.formula import ArrayFormula
TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
SUM = r'C:\Users\Jubin\Desktop\源数据\油耗分析汇总.xlsm'
CUR = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'

print('=== D17 模板 199~204 行 (F/G):')
wb = openpyxl.load_workbook(os.path.join(TPL, 'D17-4.2米-京LPW138.xlsx'), data_only=True)
wst = wb['计算模板']
for r in range(199, 205):
    print(f'  行{r}: F={wst.cell(r,6).value} G={wst.cell(r,7).value}')
wsd = wb['加油数据']
print('  D17 加油数据里 09-21 的记录:')
for i, row in enumerate(wsd.iter_rows(min_row=2, values_only=True), start=2):
    if row and row[12] is not None and str(row[12])[:10] == '2026-09-21':
        print(f'    行{i}: {row[0]} {str(row[12])[:19]} 加满[{row[8]}] ¥{row[5]}')
wb.close()

print('=== 汇总里 京LPW138 的行:')
wb = openpyxl.load_workbook(SUM, read_only=True, data_only=True)
ws = wb['Sheet1']
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if row and row[2] and str(row[2]).strip() == '京LPW138':
        print(f'  行{i}: 序号={row[0]} 时间A={str(row[5])[:19]} 时间B={str(row[6])[:19]} 金额={row[7]} 升={row[8]}')
wb.close()

print('=== 考核表里 LPW138 的行:')
wb = openpyxl.load_workbook(CUR, read_only=True, data_only=True)
ws = wb['每日油耗数据']
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if row and row[3] and str(row[3]).strip() == '京LPW138':
        print(f'  行{i}: {row[0]} F={str(row[5])[:19]} G={str(row[6])[:19]} I={row[8]} J={row[9]}')
wb.close()
