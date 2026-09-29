# -*- coding: utf-8 -*-
"""09-19 分发结果校验"""
import os, glob
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from python_calamine import CalamineWorkbook

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
EXPORT = r'C:\Users\Jubin\Desktop\数据\车辆加油_管理(2026-09-19).xlsx'

# 导出 sids
wb = CalamineWorkbook.from_path(EXPORT)
rows = wb.get_sheet_by_name(wb.sheet_names[0]).to_python()
hdr = [str(x).strip() for x in rows[0]]
exp = {}
for row in rows[1:]:
    d = dict(zip(hdr, row))
    sid = str(d.get('系统编号') or '').strip()
    if sid: exp[sid] = d
print(f'导出 {len(exp)} 条')

# 1) 加油数据: 每条 sid 都能找到, 且只在 1 个文件
found = {}
files = [fp for fp in glob.glob(os.path.join(TPL, '*.xlsx'))
         if not os.path.basename(fp).startswith(('~$', '模板'))]
for fp in files:
    fn = os.path.basename(fp)
    w = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    ws = w['加油数据']
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0] is not None and str(row[0]).strip() in exp:
            found.setdefault(str(row[0]).strip(), []).append(fn)
    w.close()
miss = [s for s in exp if s not in found]
dupd = {s: v for s, v in found.items() if len(v) > 1}
print(f'加油数据落位: {len(found)}/{len(exp)}  缺失: {miss or "无"}  重复落位: {dupd or "无"}')

# 2) 计算模板 G 与 加油数据 M 对应(排除备注排除词记录)
def gs(wb):
    for n in ('计算模板', '油耗计算'):
        if n in wb.sheetnames: return wb[n]
bad = []
checked = set()
for fp in files:
    fn = os.path.basename(fp)
    w = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    wsd, wst = w['加油数据'], gs(w)
    gset = set()
    for row in wst.iter_rows(min_row=2, values_only=True):
        if row and row[6] is not None: gset.add(str(row[6])[:19])
    for row in wsd.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None: continue
        sid = str(row[0]).strip()
        if sid in exp and sid not in checked:
            checked.add(sid)
            note = str(row[13] or '')
            kw = any(k in note for k in ('发动机','发电机','小油箱','新车'))
            mixed = '尿素' in note and '加油' in note.replace('油卡加油','')
            if not kw and not mixed and str(row[12])[:19] not in gset:
                bad.append(f'{sid} 未映射于 {fn}')
    w.close()
print(f'计算模板映射检查: {len(checked)}条  异常: {bad or "无"}')

# 3) 抽查关键行
def cell(wst, r, c):
    v = wst.cell(r, c).value
    return v.text if isinstance(v, ArrayFormula) else v
for fn, rows_chk in [('D14-4.2米-粤BT802F.xlsx', [197, 198]),
                     ('B13-9.6米-粤BNT993.xlsx', [148, 149]),
                     ('B16-9.6米-粤BNV030.xlsx', [194, 195]),
                     ('尿素B16-9.6米-粤BNV030.xlsx', [85])]:
    w = openpyxl.load_workbook(os.path.join(TPL, fn), data_only=False)
    wst = gs(w)
    for r in rows_chk:
        print(f'  {fn} 行{r}: F={cell(wst,r,6)} G={cell(wst,r,7)} I={str(cell(wst,r,9))[:70]}')
    w.close()

# 4) 考核表: 去重 & 新行
w = openpyxl.load_workbook(KAOHE, read_only=True, data_only=True)
from collections import Counter
for sn in ('每日油耗数据', '每日尿素数据'):
    ws = w[sn]
    ids = [str(row[0]).strip() for row in ws.iter_rows(min_row=2, values_only=True)
           if row and row[0] not in (None, '', '0')]
    dup = [k for k, v in Counter(ids).items() if v > 1]
    tail = [s for s in ids if s in exp]
    print(f'{sn}: 总行{len(ids)} 重复{dup or "无"} 本次新增{len(tail)}条: {sorted(tail)}')
w.close()
