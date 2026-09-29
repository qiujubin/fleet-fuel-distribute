# -*- coding: utf-8 -*-
"""修复 09-20 的处理:
1) D14: 行199(430,当天首个是->按否,F清空); 行200(434,有效是,F=09-18 23:02:36)
2) D14 跑 rebuild_formulas 重建全表公式
3) 考核表补 8 条有效是行 (7辆单车是 + D14的434累计行)
"""
import os, sys, datetime
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
WS = r'E:\WorkSpace\workbuddy\2026-09-18-16-38-15'
DT_FMT = 'yyyy-mm-dd hh:mm:ss'
sys.path.insert(0, WS)
from rebuild_formulas import rebuild

# ---- 1) D14 行199/200 修正 ----
p = os.path.join(TPL, 'D14-4.2米-粤BT802F.xlsx')
wb = openpyxl.load_workbook(p, data_only=False)
wst = wb['计算模板']
wst.cell(199, 6, None)                       # 430: 非有效, F清空
wst.cell(200, 6, datetime.datetime(2026, 9, 18, 23, 2, 36)).number_format = DT_FMT  # 434: F=上一有效是
wb.save(p); wb.close()
print('D14: 行199 F清空, 行200 F=09-18 23:02:36')

# ---- 2) D14 全表公式重建 ----
n, filled = rebuild(p)
print(f'D14 公式重建: {n}单元格, G行{filled}')

# ---- 3) 考核表补 8 行 ----
# 收集每辆车的有效是行: 模板里 G=今天记录时间 的行
targets = [
    ('260920000431', 'B13-9.6米-粤BNT993.xlsx'),
    ('260920000432', 'B14-9.6米-粤BNX948.xlsx'),
    ('260920000434', 'D14-4.2米-粤BT802F.xlsx'),
    ('260920000435', 'D19-4.2米-京N3VL85.xlsx'),
    ('260920000436', 'A1-挂车-粤BFM072.xlsx'),
    ('260920000437', 'B11-9.6米-粤BFE485.xlsx'),
    ('260920000438', 'B1-9.6米-粤ACJ959.xlsx'),
    ('260920000439', 'B12-9.6米-粤BQN066.xlsx'),
]
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

wbk = openpyxl.load_workbook(KAOHE, keep_vba=True)
for sid, fn in targets:
    fp = os.path.join(TPL, fn)
    wbv = openpyxl.load_workbook(fp, data_only=True)
    wstv = wbv['计算模板']; wsdv = wbv['加油数据']
    plate = str(wstv.cell(2, 4).value or '').strip()
    type_v = str(wstv.cell(2, 5).value or '')
    # 找记录行
    rec = None
    for r in range(2, wsdv.max_row + 1):
        if str(wsdv.cell(r, 1).value or '').strip() == sid:
            rec = {'row': r, 'mile': num(wsdv.cell(r, 4).value), 'gps': num(wsdv.cell(r, 5).value),
                   'amt': num(wsdv.cell(r, 6).value), 'vol': num(wsdv.cell(r, 7).value),
                   'price': num(wsdv.cell(r, 8).value), 'person': wsdv.cell(r, 12).value,
                   'time': wsdv.cell(r, 13).value, 'note': str(wsdv.cell(r, 14) or '').strip()}
            break
    assert rec, (sid, fn)
    # 模板里该记录的行 (G==登记时间)
    tr = None
    for r in range(2, wstv.max_row + 1):
        if wstv.cell(r, 7).value and str(wstv.cell(r, 7).value)[:19] == str(rec['time'])[:19]:
            tr = r; break
    assert tr, (sid, '模板行未找到')
    F_time = wstv.cell(tr, 6).value
    # 周期行: G==F 的行 +1 .. tr-1
    cycle = []
    if F_time:
        anchor = None
        for r in range(2, wstv.max_row + 1):
            if wstv.cell(r, 7).value and str(wstv.cell(r, 7).value)[:19] == str(F_time)[:19]:
                anchor = r; break
        if anchor: cycle = list(range(anchor + 1, tr))
    amt, vol = rec['amt'], rec['vol']
    for x in cycle:
        gt = wstv.cell(x, 7).value
        for r2 in range(2, wsdv.max_row + 1):
            if wsdv.cell(r2, 13).value and str(wsdv.cell(r2, 13).value)[:19] == str(gt)[:19]:
                a, v = num(wsdv.cell(r2, 6).value), num(wsdv.cell(r2, 7).value)
                amt = (amt or 0) + (a or 0); vol = (vol or 0) + (v or 0)
                break
    recF = None
    for r2 in range(2, wsdv.max_row + 1):
        if wsdv.cell(r2, 13).value and str(wsdv.cell(r2, 13).value)[:19] == str(F_time)[:19]:
            recF = (num(wsdv.cell(r2, 4).value), num(wsdv.cell(r2, 5).value)); break
    L = recF[1] if recF else None   # GPS开始
    M = recF[0] if recF else None   # 仪表开始
    N, O = rec['gps'], rec['mile']
    P = round(N - L, 4) if (L is not None and N is not None) else None
    Q = round(O - M, 4) if (M is not None and O is not None) else None
    Rv = round(vol / P * 100, 6) if P else None
    Sv = round(amt / P, 6) if P else None
    Tv = round(vol / Q * 100, 6) if Q else None
    Uv = round(amt / Q, 10) if Q else None
    Vv = round((Rv + Tv) / 2, 10) if (Rv is not None and Tv is not None) else None
    sheet = '每日尿素数据' if fn.startswith('尿素') else '每日油耗数据'
    ws = wbk[sheet]
    exists = any(str(ws.cell(x, 1).value or '').strip() == sid for x in range(2, ws.max_row + 1))
    if exists:
        print(f'跳过(已在考核): {sid}'); continue
    nr = 1
    for x in range(1, ws.max_row + 1):
        if any(ws.cell(x, c).value not in (None, '') for c in range(1, 25)): nr = x
    nr += 1
    rowvals = [sid, tr - 1, rec['person'], plate, type_v, F_time, rec['time'], '是',
               amt, vol, rec['price'], L, M, N, O, P, Q, Rv, Sv, Tv, Uv, Vv,
               rec['time'].date(), rec['note'] or 0]
    for c, v in enumerate(rowvals, 1):
        cell = ws.cell(nr, c, v)
        if c in (6, 7): cell.number_format = DT_FMT
        if c == 23: cell.number_format = 'yyyy-mm-dd'
    print(f'{sid} {plate} -> {sheet} 第{nr}行 (I={amt}, J={vol}, 周期{len(cycle)}行)')
    wbv.close()
wbk.save(KAOHE)
print('考核表已保存')
import subprocess
subprocess.run([sys.executable, os.path.join(WS, 'final_repair.py')], check=False)
