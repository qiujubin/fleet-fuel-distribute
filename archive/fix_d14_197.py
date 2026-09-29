# -*- coding: utf-8 -*-
"""修 D14 行197(408, 09-18 12:14): F 应清空(当天最后"是"是 198 的 415)"""
import sys, os
sys.path.insert(0, r'E:\WorkSpace\workbuddy\2026-09-18-16-38-15')
import openpyxl
from rebuild_formulas import rebuild

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
D14 = os.path.join(TPL, 'D14-4.2米-粤BT802F.xlsx')

wb = openpyxl.load_workbook(D14, data_only=False)
wst = wb['计算模板']
print('改前: F197=', wst.cell(197, 6).value, '| F198=', wst.cell(198, 6).value)
assert wst.cell(197, 6).value is not None and wst.cell(198, 6).value is not None
wst.cell(197, 6).value = None
wb.save(D14); wb.close()
n, filled = rebuild(D14)
print(f'rebuild: {n} 公式 / {filled} 行G')

import pythoncom
import win32com.client as wc
pythoncom.CoInitialize()
app = None
try:
    app = wc.DispatchEx('Ket.Application')
    app.Visible = False; app.DisplayAlerts = False
    wb2 = app.Workbooks.Open(D14, 0, False)
    ws = wb2.Worksheets('计算模板')
    app.Calculate()
    print('重算后 197/198/201 行:')
    for r in (197, 198, 201):
        fv = ws.Range(ws.Cells(r, 6), ws.Cells(r, 10)).Value2
        fv = fv[0] if isinstance(fv[0], tuple) else fv
        print(f'  行{r}: F={fv[0]!r} G={fv[1]!r} I={fv[3]!r}')
    wb2.Save(); wb2.Close(False)
finally:
    if app is not None:
        try: app.Quit()
        except Exception: pass
    pythoncom.CoUninitialize()
print('完成')
