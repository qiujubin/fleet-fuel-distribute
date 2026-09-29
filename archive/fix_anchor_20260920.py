# -*- coding: utf-8 -*-
"""综合修复: 今天(09-19下午/09-20)新增的"是"行 F 锚点全部重算 + 公式重建 + 考核行重写
F 正确值 = 模板中本行之前最后一个"是"行的 G (从加油数据反查是否加满, 不依赖H缓存)"""
import os, sys, datetime
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
WS = r'E:\WorkSpace\workbuddy\2026-09-18-16-38-15'
sys.path.insert(0, WS)
from rebuild_formulas import rebuild

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

# 今天新分发的"是"记录 (sid -> 模板文件)
DISTRIBUTED_SHI = [
    ('260919000427', 'D13-4.2米-粤BT3S56.xlsx'),
    ('260919000429', 'B11-9.6米-粤BFE485.xlsx'),
    ('260920000431', 'D17-4.2米-京LPW138.xlsx'),
    ('260920000432', 'B14-9.6米-粤BNX948.xlsx'),
    ('260920000435', 'D19-4.2米-京N3VL85.xlsx'),
    ('260920000436', 'A1-挂车-粤BFM072.xlsx'),
    ('260920000437', 'D7-4.2米-粤B61HY1.xlsx'),
    ('260920000438', 'D18-4.2米-京HPN783.xlsx'),
    ('260920000439', 'B12-9.6米-粤BQN066.xlsx'),
]
# 按文件分组(一个文件可能多行)
from collections import defaultdict
by_file = defaultdict(list)
for sid, fn in DISTRIBUTED_SHI:
    by_file[fn].append(sid)

kao_updates = []   # (sid, fn) 考核行待重写
for fn, sids in by_file.items():
    fp = os.path.join(TPL, fn)
    wb = openpyxl.load_workbook(fp, data_only=False)
    wst = wb['计算模板']; wsd = wb['加油数据']
    # 加油数据: 时间->是否加满
    full_map = {}
    for r in range(2, wsd.max_row + 1):
        mt = wsd.cell(r, 13).value
        if mt is not None:
            full_map[str(mt)[:19]] = str(wsd.cell(r, 9).value or '').strip()
    # 模板 G 行清单
    g_rows = []
    for r in range(2, wst.max_row + 1):
        g = wst.cell(r, 7).value
        if g is not None:
            g_rows.append((r, str(g)[:19]))
    # 重算今天新"是"行的 F
    fixed_any = False
    for r, g19 in g_rows:
        if g19 < '2026-09-19' and not any(g19 == str(datetime.datetime(2026,9,19,15,32,46))[:19] for _ in [0]):
            pass
    # 简化: 找到目标 sid 对应的模板行, 重算 F
    for sid in sids:
        # 该记录的登记时间
        mt = None
        for r in range(2, wsd.max_row + 1):
            if str(wsd.cell(r, 1).value or '').strip() == sid:
                mt = str(wsd.cell(r, 13).value)[:19]; break
        tr = g_rows_map = None
        for r, g19 in g_rows:
            if g19 == mt:
                tr = r; break
        assert tr, (fn, sid, mt)
        # 上一个"是"行 (模板顺序, 本行之前)
        prev_shi_G = None
        for r2, g19 in g_rows:
            if r2 >= tr: break
            if full_map.get(g19) == '是': prev_shi_G = g19
        cur_F = wst.cell(tr, 6).value
        cur_F19 = str(cur_F)[:19] if cur_F is not None else None
        if cur_F19 != prev_shi_G:
            if prev_shi_G is None:
                wst.cell(tr, 6).value = None
            else:
                wst.cell(tr, 6, datetime.datetime.strptime(prev_shi_G, '%Y-%m-%d %H:%M:%S')).number_format = 'yyyy-mm-dd hh:mm:ss'
            print(f'{fn} {sid} 行{tr}: F {cur_F19} -> {prev_shi_G}')
            fixed_any = True
        kao_updates.append((sid, fn, tr))
    if fixed_any:
        wb.save(fp)
        n, filled = rebuild(fp)
        print(f'{fn}: 公式重建 {n}单元格/{filled}行G')
    else:
        print(f'{fn}: F 已正确, 无需改')
    wb.close()

# ---- 考核行重写 ----
wbk = openpyxl.load_workbook(KAOHE, keep_vba=True)
for sid, fn, tr in kao_updates:
    fp = os.path.join(TPL, fn)
    wbv = openpyxl.load_workbook(fp, data_only=True)
    wstv = wbv['计算模板']; wsdv = wbv['加油数据']
    plate = str(wstv.cell(2, 4).value or '').strip()
    type_v = str(wstv.cell(2, 5).value or '')
    rec = None
    for r in range(2, wsdv.max_row + 1):
        if str(wsdv.cell(r, 1).value or '').strip() == sid:
            rec = {'mile': num(wsdv.cell(r, 4).value), 'gps': num(wsdv.cell(r, 5).value),
                   'amt': num(wsdv.cell(r, 6).value), 'vol': num(wsdv.cell(r, 7).value),
                   'price': num(wsdv.cell(r, 8).value), 'person': wsdv.cell(r, 12).value,
                   'time': wsdv.cell(r, 13).value, 'note': str(wsdv.cell(r, 14) or '').strip()}
            break
    F_time = wstv.cell(tr, 6).value
    cycle = []
    if F_time:
        anchor = None
        for r in range(2, wstv.max_row + 1):
            if wstv.cell(r, 7).value and str(wstv.cell(r, 7).value)[:19] == str(F_time)[:19]:
                anchor = r; break
        if anchor: cycle = list(range(anchor + 1, tr))
    amt, vol = rec['amt'], rec['vol']
    for x in cycle:
        gt = wstv.cell(x, 7).value
        if gt is None: continue
        for r2 in range(2, wsdv.max_row + 1):
            if wsdv.cell(r2, 13).value and str(wsdv.cell(r2, 13).value)[:19] == str(gt)[:19]:
                a, v = num(wsdv.cell(r2, 6).value), num(wsdv.cell(r2, 7).value)
                amt = (amt or 0) + (a or 0); vol = (vol or 0) + (v or 0)
                break
    recF = None
    if F_time:
        for r2 in range(2, wsdv.max_row + 1):
            if wsdv.cell(r2, 13).value and str(wsdv.cell(r2, 13).value)[:19] == str(F_time)[:19]:
                recF = (num(wsdv.cell(r2, 4).value), num(wsdv.cell(r2, 5).value)); break
    L = recF[1] if recF else None; M = recF[0] if recF else None
    N, O = rec['gps'], rec['mile']
    P = round(N - L, 4) if (L is not None and N is not None) else None
    Q = round(O - M, 4) if (M is not None and O is not None) else None
    Rv = round(vol / P * 100, 6) if P else None
    Sv = round(amt / P, 6) if P else None
    Tv = round(vol / Q * 100, 6) if Q else None
    Uv = round(amt / Q, 10) if Q else None
    Vv = round((Rv + Tv) / 2, 10) if (Rv is not None and Tv is not None) else None
    ws = wbk['每日油耗数据']
    trow = None
    for x in range(2, ws.max_row + 1):
        if str(ws.cell(x, 1).value or '').strip() == sid:
            trow = x; break
    assert trow, (sid, '考核行未找到')
    rowvals = {2: tr - 1, 6: F_time, 9: amt, 10: vol, 12: L, 13: M, 14: N, 15: O,
               16: P, 17: Q, 18: Rv, 19: Sv, 20: Tv, 21: Uv, 22: Vv}
    for c, v in rowvals.items():
        cell = ws.cell(trow, c, v)
        if c == 6: cell.number_format = 'yyyy-mm-dd hh:mm:ss'
    print(f'考核行重写 {sid} (第{trow}行): I={amt} J={vol} P={P} Q={Q}')
    wbv.close()
wbk.save(KAOHE)
print('考核表已保存')
import subprocess
subprocess.run([sys.executable, os.path.join(WS, 'final_repair.py')], check=False)
print('全部完成')
