# -*- coding: utf-8 -*-
import openpyxl, os
from openpyxl.worksheet.formula import ArrayFormula
TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
CUR = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'

wb = openpyxl.load_workbook(os.path.join(TPL, 'D15-4.2米-粤BY2J27.xlsx'), data_only=False)
wst = wb['计算模板']
wsd = wb['加油数据']
def txt(c):
    v = wst[c].value
    return v.text if isinstance(v, ArrayFormula) else v
r = 145
print('D15 模板145: F=', txt(f'F{r}'), 'G=', txt(f'G{r}'))
print('  I=', str(txt(f'I{r}'))[:80])
for rr in range(2, wsd.max_row + 1):
    if str(wsd.cell(rr, 1).value or '').strip() == '260920000443':
        print('加油数据行', rr, ':', str(wsd.cell(rr, 1).value), str(wsd.cell(rr, 13).value)[:19],
              '加满', wsd.cell(rr, 9).value, '| A列类型:', type(wsd.cell(rr, 1).value).__name__)
wb.close()

wb = openpyxl.load_workbook(CUR, read_only=True, data_only=True)
ws = wb['每日油耗数据']
for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    if row and str(row[0] or '').strip() == '260920000443':
        print('考核行', i, ': 车牌', row[3], '| F=', row[5], 'G=', row[6],
              '| I=', row[8], 'J=', row[9], '| 列数', len([c for c in row if c is not None]))
        print('   A列类型:', type(row[0]).__name__)
cnt = sum(1 for row in ws.iter_rows(min_row=2, values_only=True)
          if row and str(row[0] or '').strip() == '260920000443')
print('考核表 443 出现次数:', cnt)
wb.close()
