# -*- coding: utf-8 -*-
"""修复 09-19 首次分发的三个问题:
1. D14/B13 两条是记录覆盖同一行 -> 重排两行
2. B16 的 410 (单价2.6=尿素) 从油文件搬到尿素文件
3. 考核表序号修正 + 410 考核行删除/重建
"""
import os, datetime
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
DT_FMT = 'yyyy-mm-dd hh:mm:ss'

def shift_formula(text, src, dst):
    import re
    return re.sub(r'([A-Z])%d\b' % src, lambda m: m.group(1) + str(dst), text)

def time_of_sid(wsd, sid):
    for r in range(2, wsd.max_row + 1):
        if str(wsd.cell(r, 1).value or '').strip() == sid:
            return r, wsd.cell(r, 13).value
    return None, None

def base_ij(r):
    i = f'=XLOOKUP(D{r}&G{r},加油数据!B:B&加油数据!M:M,加油数据!F:F,0)'
    j = f'=XLOOKUP(D{r}&G{r},加油数据!B:B&加油数据!M:M,加油数据!G:G,0)'
    return ArrayFormula(f'I{r}', i), ArrayFormula(f'J{r}', j)

def get_sheets(path):
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    def gs(wb):
        for n in ('计算模板', '油耗计算'):
            if n in wb.sheetnames: return wb[n]
    return wb_f, wb_v, gs(wb_f), gs(wb_v), wb_f['加油数据'], wb_v['加油数据']

# ============ A. B13-粤BNT993: 行148/149 重排 ============
p = os.path.join(TPL, 'B13-9.6米-粤BNT993.xlsx')
wb_f, wb_v, wst_f, wst_v, wsd_f, wsd_v = get_sheets(p)
_, t409 = time_of_sid(wsd_v, '260918000409')
_, t424 = time_of_sid(wsd_v, '260919000424')
assert t409 and t424
# 行148 = 409(是): F 保持 09-17 17:27:56, G 改为 409 时间, I/J base
wst_f.cell(148, 7, t409).number_format = DT_FMT
i, j = base_ij(148)
wst_f.cell(148, 9, i); wst_f.cell(148, 10, j)
# 行149 = 424(是): F = 409 时间, G = 424 时间, I/J base
wst_f.cell(149, 6, t409).number_format = DT_FMT
wst_f.cell(149, 7, t424).number_format = DT_FMT
i, j = base_ij(149)
wst_f.cell(149, 9, i); wst_f.cell(149, 10, j)
wb_f.save(p)
print('A. B13 修复: 行148->409, 行149->424')

# ============ B. D14-粤BT802F: 行197/198 重排 ============
p = os.path.join(TPL, 'D14-4.2米-粤BT802F.xlsx')
wb_f, wb_v, wst_f, wst_v, wsd_f, wsd_v = get_sheets(p)
_, t408 = time_of_sid(wsd_v, '260918000408')
_, t415 = time_of_sid(wsd_v, '260918000415')
assert t408 and t415
wst_f.cell(197, 7, t408).number_format = DT_FMT
i, j = base_ij(197)
wst_f.cell(197, 9, i); wst_f.cell(197, 10, j)
wst_f.cell(198, 6, t408).number_format = DT_FMT
wst_f.cell(198, 7, t415).number_format = DT_FMT
i, j = base_ij(198)
wst_f.cell(198, 9, i); wst_f.cell(198, 10, j)
wb_f.save(p)
print('B. D14 修复: 行197->408, 行198->415')

# ============ C. B16-粤BNV030: 410 搬走, 411 挪回行194 ============
p = os.path.join(TPL, 'B16-9.6米-粤BNV030.xlsx')
wb_f, wb_v, wst_f, wst_v, wsd_f, wsd_v = get_sheets(p)
r410, t410 = time_of_sid(wsd_v, '260918000410')
r411, t411 = time_of_sid(wsd_v, '260918000411')
assert r410 and r411
row410 = [wsd_f.cell(r410, c).value for c in range(1, 15)]
wsd_f.delete_rows(r410, 1)          # 删 410, 411 上移
# 计算模板: 行194 <- 411(否): 恢复 I/J 骨架(从196行复制平移), F 清空, G=411
src = 196
for c in (9, 10):
    v = wst_f.cell(src, c).value
    col = get_column_letter(c)
    if isinstance(v, ArrayFormula):
        wst_f.cell(194, c, ArrayFormula(f'{col}194', shift_formula(v.text, src, 194)))
wst_f.cell(194, 6, None)
wst_f.cell(194, 7, t411).number_format = DT_FMT
# 行195 恢复空(骨架)
wst_f.cell(195, 7, None)
wst_f.cell(195, 6, None)
wb_f.save(p)
print('C. B16 修复: 410 已从油文件移除, 411 挪到行194, 410数据已暂存')

# ============ D. 尿素B16: 加入 410 ============
p = os.path.join(TPL, '尿素B16-9.6米-粤BNV030.xlsx')
wb_f, wb_v, wst_f, wst_v, wsd_f, wsd_v = get_sheets(p)
# 加油数据追加
sids = set(str(wsd_v.cell(r, 1).value or '').strip() for r in range(2, wsd_v.max_row+1))
assert '260918000410' not in sids
lr = 1
for r in range(1, wsd_f.max_row + 1):
    if any(wsd_f.cell(r, c).value not in (None, '') for c in range(1, 15)): lr = r
nr = lr + 1
for c, v in enumerate(row410, 1):
    wsd_f.cell(nr, c, v)
wsd_f.cell(nr, 3).number_format = 'yyyy-mm-dd'
wsd_f.cell(nr, 13).number_format = DT_FMT
# 计算模板: 下一个空行, 是行
g_rows = [(r, wst_v.cell(r, 7).value, wst_v.cell(r, 8).value)
          for r in range(2, wst_v.max_row + 1) if wst_v.cell(r, 7).value is not None]
last_filled = g_rows[-1][0] if g_rows else 1
r = last_filled + 1
wst_f.cell(r, 7, t410).number_format = DT_FMT
prev_shi_G, cycle_start = None, 2
for rr, g, h in g_rows:
    if h == '是': prev_shi_G, cycle_start = g, rr + 1
if prev_shi_G is not None:
    wst_f.cell(r, 6, prev_shi_G).number_format = DT_FMT
    off_i = ''.join(f'+I{x}' for x in range(cycle_start, r))
    off_j = ''.join(f'+J{x}' for x in range(cycle_start, r))
else:
    first_t = wsd_v.cell(2, 13).value
    wst_f.cell(r, 6, first_t).number_format = DT_FMT
    off_i = off_j = ''
base_i = f'=XLOOKUP(D{r}&G{r},加油数据!B:B&加油数据!M:M,加油数据!F:F,0)'
base_j = f'=XLOOKUP(D{r}&G{r},加油数据!B:B&加油数据!M:M,加油数据!G:G,0)'
wst_f.cell(r, 9, ArrayFormula(f'I{r}', base_i + off_i))
wst_f.cell(r, 10, ArrayFormula(f'J{r}', base_j + off_j))
print(f'D. 尿素B16: 410 加入加油数据第{nr}行, 计算模板第{r}行[是] F={prev_shi_G} 偏移I={off_i or "无"}')
wb_f.save(p)

# ============ E. 考核表修正 ============
wbk = openpyxl.load_workbook(KAOHE, keep_vba=True)
w1 = wbk['每日油耗数据']; w2 = wbk['每日尿素数据']

def find_row(ws, sid):
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or '').strip() == sid: return r
    return None

# 修正序号
for sid, xh in [('260918000409', 147), ('260919000424', 148), ('260918000408', 196), ('260918000415', 197)]:
    rr = find_row(w1, sid)
    if rr:
        old = w1.cell(rr, 2).value
        w1.cell(rr, 2, xh)
        print(f'E1. 考核表 {sid} 序号 {old} -> {xh}')
# 删除 410 在每日油耗数据的行
rr = find_row(w1, '260918000410')
if rr:
    w1.delete_rows(rr, 1)
    print('E2. 考核表: 410 已从每日油耗数据删除')
# 410 加入每日尿素数据(计算值)
amt, vol = float(row410[5] or 0), float(row410[6] or 0)
plate_v = '粤BNV030'
# 尿素B16 周期内其他行金额
g_rows2 = [(r, wst_v.cell(r, 7).value, wst_v.cell(r, 8).value)
           for r in range(2, wst_v.max_row + 1) if wst_v.cell(r, 7).value is not None]
prev_shi_G, cycle_start = None, 2
for rr2, g, h in g_rows2:
    if h == '是' and str(g)[:19] < str(t410)[:19]:
        prev_shi_G, cycle_start = g, rr2 + 1
r_new = last_filled + 1
for x in range(cycle_start, r_new):
    gt = wst_v.cell(x, 7).value
    if gt is None: continue
    for r3 in range(2, wsd_v.max_row + 1):
        if str(wsd_v.cell(r3, 13).value or '')[:19] == str(gt)[:19]:
            try: amt += float(wsd_v.cell(r3, 6).value or 0)
            except: pass
            try: vol += float(wsd_v.cell(r3, 7).value or 0)
            except: pass
            break
recF = None
for r3 in range(2, wsd_v.max_row + 1):
    if wsd_v.cell(r3, 13).value and str(wsd_v.cell(r3, 13).value)[:19] == str(prev_shi_G)[:19]:
        recF = (num4 := wsd_v.cell(r3, 5).value, wsd_v.cell(r3, 4).value)
        break
L = float(recF[0]) if recF and recF[0] else None
M = float(recF[1]) if recF and recF[1] else None
N, O = float(row410[4] or 0), float(row410[3] or 0)
P = round(N - L, 4) if L is not None else None
Q = round(O - M, 4) if M is not None else None
Rv = round(vol / P * 100, 6) if P else None
Sv = round(amt / P, 6) if P else None
Tv = round(vol / Q * 100, 6) if Q else None
lr2 = 1
for r in range(1, w2.max_row + 1):
    if any(w2.cell(r, c).value not in (None, '') for c in range(1, 21)): lr2 = r
nr2 = lr2 + 1
rowvals = ['260918000410', r_new - 1, row410[11], plate_v, '9.6米',
           prev_shi_G, t410, '是', amt, vol, row410[7], L, M, N, O, P, Q, Rv, Sv, Tv]
for c, v in enumerate(rowvals, 1):
    cell = w2.cell(nr2, c, v)
    if c in (6, 7): cell.number_format = DT_FMT
print(f'E3. 考核表: 410 加入每日尿素数据第{nr2}行 (I={amt}, J={vol})')
wbk.save(KAOHE)
print('考核表已保存')
