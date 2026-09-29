# -*- coding: utf-8 -*-
"""修 D14 粤BT802F 09-20 多"是"回溯: 434 降级为否, 448 升为有效是(当天最后一个"是")
步骤: ①改 F 值 ②rebuild 公式 ③COM 重算读实时值 ④替换考核表 434 行为 448"""
import sys, os, datetime
sys.path.insert(0, r'E:\WorkSpace\workbuddy\2026-09-18-16-38-15')
import openpyxl
from rebuild_formulas import rebuild

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
D14 = os.path.join(TPL, 'D14-4.2米-粤BT802F.xlsx')

# ---- 1) 改 F 值: 行200(434) 清空, 行201(448) = 09-18 23:02:36(415的G) ----
wb = openpyxl.load_workbook(D14, data_only=False)
wst = wb['计算模板']
print('改前: F200=', wst.cell(200, 6).value, ' F201=', wst.cell(201, 6).value)
wst.cell(200, 6).value = None
wst.cell(201, 6).value = datetime.datetime(2026, 9, 18, 23, 2, 36)
wb.save(D14)
wb.close()
print('F 值已改: 行200 清空, 行201 = 2026-09-18 23:02:36')

# ---- 2) 重建公式 ----
n, filled = rebuild(D14)
print(f'rebuild: {n} 公式单元格, {filled} 行G')

# ---- 3) COM 重算并读 201 行实时值 ----
import pythoncom
import win32com.client as wc
pythoncom.CoInitialize()
app = None
vals201 = None
try:
    app = wc.DispatchEx('Ket.Application')
    app.Visible = False; app.DisplayAlerts = False
    wb2 = app.Workbooks.Open(D14, 0, False)
    assert not wb2.ReadOnly
    ws2 = wb2.Worksheets('计算模板')
    app.Calculate()
    rng = ws2.Range(ws2.Cells(201, 6), ws2.Cells(201, 22)).Value2   # F..V
    vals201 = rng[0] if isinstance(rng[0], tuple) else rng
    print('201行实时值 F..V:', [round(v, 4) if isinstance(v, float) else v for v in vals201])
    i200 = ws2.Cells(200, 9).Value2
    print('200行(434) I =', i200)
    wb2.Save(); wb2.Close(False)
finally:
    if app is not None:
        try: app.Quit()
        except Exception: pass
    pythoncom.CoUninitialize()

# ---- 4) 考核表: 替换 434 行为 448 (必须用 COM, openpyxl 保存会剥掉 WPS 宏部件) ----
from openpyxl.utils.datetime import from_excel
pythoncom.CoInitialize()
app = None
try:
    app = wc.DispatchEx('Ket.Application')
    app.Visible = False; app.DisplayAlerts = False
    wbk = app.Workbooks.Open(KAOHE, 0, False)
    assert not wbk.ReadOnly
    ws = wbk.Worksheets('每日油耗数据')
    target = None
    for i in range(2, 6000):
        v = ws.Cells(i, 1).Value2
        if v is not None and str(v).strip() in ('260920000434', '260920000434.0'):
            target = i; break
    assert target, '找不到 434 的考核行'
    print('找到 434 考核行:', target)
    def n(v):
        try: return float(v)
        except (TypeError, ValueError): return None
    F_s, G_s = vals201[0], vals201[1]          # 序列数
    person = ws.Cells(target, 3).Value2
    plate = ws.Cells(target, 4).Value2
    ptype = ws.Cells(target, 5).Value2
    seq = ws.Cells(target, 2).Value2
    rowvals = ['260920000448', seq, person, plate, ptype,
               F_s, G_s, '是',
               n(vals201[3]), n(vals201[4]), 7.95,
               n(vals201[6]), n(vals201[7]), n(vals201[8]), n(vals201[9]),
               n(vals201[10]), n(vals201[11]), n(vals201[12]), n(vals201[13]), n(vals201[14]),
               n(vals201[15]), n(vals201[16]),
               float(int(G_s)), 0]
    ws.Cells(target, 1).NumberFormat = '@'
    ws.Range(ws.Cells(target, 1), ws.Cells(target, 24)).Value = tuple([tuple(rowvals)])
    ws.Cells(target, 6).NumberFormat = 'yyyy-mm-dd hh:mm:ss'
    ws.Cells(target, 7).NumberFormat = 'yyyy-mm-dd hh:mm:ss'
    ws.Cells(target, 23).NumberFormat = 'yyyy-mm-dd'
    wbk.Save()
    print(f'考核行 {target} 已替换为 448: I={n(vals201[3])} J={n(vals201[4])}')
    wbk.Close(False)
finally:
    if app is not None:
        try: app.Quit()
        except Exception: pass
    pythoncom.CoUninitialize()
print('完成')
