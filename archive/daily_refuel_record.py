# -*- coding: utf-8 -*-
"""每日加油记录 sheet 刷新: 取 n-1(数据中最大登记日期-1) 的加油记录
规则: 排除尿素及错标尿素(单价<5); 排除词记录"是"改"否";
      同车牌同日多条 -> 有"是"留最后一个"是", 全"否"留最后一个。整表替换。"""
import os, glob, sys, datetime
import openpyxl

TPL = r'C:\Users\Jubin\Desktop\车辆维修，油耗\车辆油耗计算模板'
KAOHE = r'C:\Users\Jubin\Desktop\车辆维修，油耗\油耗考核10月_宏.xlsm'
RESTORE = r'E:\WorkSpace\workbuddy\2026-09-18-16-38-15\restore_wps_parts.py'
EXCLUDE_KW = ['发动机', '发电机', '小油箱', '新车']
DT_FMT, D_FMT = 'yyyy-mm-dd hh:mm:ss', 'yyyy-mm-dd'

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return v

# 1) 从各车辆文件收集 n-1 的候选记录
def to_date(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    s = str(v).strip()
    try: return datetime.datetime.strptime(s, '%Y-%m-%d').date()
    except ValueError: return None   # 序列数等无法解析的跳过

best_date = None
cands = []
for fp in sorted(glob.glob(os.path.join(TPL, '*.xlsx'))):
    fn = os.path.basename(fp)
    if fn.startswith(('~$', '模板', '尿素')): continue
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    ws = wb['加油数据']
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None: continue
        d = to_date(row[2])
        if d is None: continue
        d10 = d.isoformat()
        if best_date is None or d10 > best_date: best_date = d10
        cands.append(row)
    wb.close()
n1 = (datetime.datetime.strptime(best_date, '%Y-%m-%d') - datetime.timedelta(days=1)).date()
print(f'数据最大日期 n={best_date} -> 取 n-1 = {n1}')

pool = []
for row in cands:
    d = to_date(row[2])
    if d is None or d != n1:
        continue
    typ = str(row[9] or '').strip()
    price = num(row[7])
    note = str(row[13] or '').strip()
    if '尿素' in typ: continue                     # 尿素不要
    if isinstance(price, float) and 0 < price < 5: continue  # 错标成加油的尿素
    full = str(row[8] or '').strip()
    if any(k in note for k in EXCLUDE_KW): full = '否'   # 排除词: 是改否
    pool.append({'row': row, 'time': row[12], 'plate': str(row[1] or '').strip().upper(), 'full': full})

# 2) 同车牌同日: 有"是"留最后一个"是", 全"否"留最后一个
by_plate = {}
for p in pool:
    by_plate.setdefault(p['plate'], []).append(p)
keep = []
for plate, items in by_plate.items():
    items.sort(key=lambda x: str(x['time'])[:19])
    shi = [x for x in items if x['full'] == '是']
    keep.append(shi[-1] if shi else items[-1])
keep.sort(key=lambda x: str(x['time'])[:19])
print(f'n-1({n1}) 候选{len(pool)}条, 保留{len(keep)}条:')

# 3) 重写 每日加油记录
wbk = openpyxl.load_workbook(KAOHE, keep_vba=True)
ws = wbk['每日加油记录']
# 清空旧行
last = 1
for r in range(1, ws.max_row + 1):
    if any(ws.cell(r, c).value not in (None, '') for c in range(1, 15)): last = r
for r in range(2, last + 1):
    for c in range(1, 15):
        ws.cell(r, c, None)
# 写新行
for i, p in enumerate(keep, start=2):
    row = p['row']
    for c, v in enumerate(row[:14], start=1):
        cell = ws.cell(i, c, v)
    ws.cell(i, 3).number_format = D_FMT
    ws.cell(i, 13).number_format = DT_FMT
    print(f"  行{i}: {row[0]} {row[1]} {str(row[2])[:10]} 加满[{p['full']}] ¥{row[5]} {row[6]}L 备注[{str(row[13] or '') or '-'}]")
wbk.save(KAOHE)
print('每日加油记录已刷新')
# 4) 恢复 WPS 宏部件
import subprocess
subprocess.run([sys.executable, RESTORE], check=False)
print('WPS部件已恢复')
